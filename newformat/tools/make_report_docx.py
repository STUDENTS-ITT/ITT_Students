"""Сборка отчёта Word по датасету dataset_1786981733."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "docs" / "figures"
OUT = ROOT / "docs" / "Отчет.docx"
OUT_ALT = ROOT / "docs" / "Отчет_способы_уменьшить_ошибки.docx"
DATA_URL = "https://drive.google.com/file/d/1pox9GSTWwSeGekAxN1-u2y8ycIR2mtac/view?usp=sharing"


def set_run_font(run, name="Times New Roman", size=12, bold=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size=16 if level == 1 else 14, bold=True)
    return p


def add_para(doc, text, *, bold=False, italic=False, size=12, space_after=8, first_line=True):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if first_line:
        pf.first_line_indent = Cm(1.25)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    run.italic = italic
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    run = p.add_run(text)
    set_run_font(run, size=12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    return p


def add_link(doc, url):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run(url)
    set_run_font(run, size=12)
    run.font.color.rgb = RGBColor(0, 0, 180)
    run.underline = True
    r_id = p.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    hyperlink.append(run._element)
    p._p.append(hyperlink)
    return p


def shade_header(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "D9E2F3")
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        set_run_font(run, size=11, bold=True)
        shade_header(cell)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_i else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(str(val))
            set_run_font(run, size=11)
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return table


def add_picture(doc, path, caption, width_cm=16.5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run()
    run.add_picture(str(path), width=Cm(width_cm))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.first_line_indent = Cm(0)
    cap.paragraph_format.space_after = Pt(12)
    run = cap.add_run(caption)
    set_run_font(run, size=11, bold=False)
    run.italic = True


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(1.5)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(6)
    r = title.add_run("ОТЧЁТ")
    set_run_font(r, size=18, bold=True)

    st = doc.add_paragraph()
    st.alignment = WD_ALIGN_PARAGRAPH.CENTER
    st.paragraph_format.space_after = Pt(6)
    r = st.add_run("по обработке навигационных данных (второй датасет)")
    set_run_font(r, size=16, bold=True)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(18)
    r = sub.add_run("Комплексирование БИНС и СНС, фильтр Калмана 15-го порядка")
    set_run_font(r, size=13)
    r.italic = True

    add_para(
        doc,
        "Программа imitator (копия newformat) выполняет счисление БИНС с коррекцией "
        "фильтром Калмана. Решение — tools/result.txt, эталон СНС — tools/reference.txt, "
        "вектор ошибок фильтра — tools/errors.txt. Ошибка на графиках: Калман − GPS "
        "(для углов — Калман − истинные углы из angle.data).",
        first_line=False,
    )

    add_heading(doc, "1. Исходные данные", 1)
    add_para(doc, "Навигационные данные взяты с Google Drive:", first_line=False)
    add_link(doc, DATA_URL)
    add_para(
        doc,
        "Архив dataset_1786981733. Файлы скопированы в newformat/data/raw/. "
        "Использованы три файла измерений:",
        first_line=False,
    )
    add_table(
        doc,
        ["Файл", "Что берём", "Роль в программе"],
        [
            [
                "imu0/imu.data",
                "гироскопы wx, wy, wz (рад/с); акселерометры ax, ay, az (м/с²)",
                "счисление БИНС и выставка",
            ],
            [
                "gps0/gps.data",
                "широта, долгота, высота; скорости vx, vy, vz",
                "эталон положения и скорости",
            ],
            [
                "state_ground_truth0/angle.data",
                "истинные углы roll, pitch, yaw (рад)",
                "эталон ориентации",
            ],
        ],
        [4.5, 7.0, 5.0],
    )
    add_para(
        doc,
        "Время в файлах записано как ЧЧ:ММ:СС.дробь (например 00:00:00.006000128). "
        "Длительность записи около 704 с. Маршрут — квадрат с набором высоты "
        "примерно от 124 м до 273 м. Стартовые координаты GPS: широта 47.64147°, "
        "долгота −122.14017°, высота 123.64 м. Время автономной выставки — 120 с.",
    )
    add_para(
        doc,
        "Чтение изменено только под этот формат: время переводится в секунды, "
        "IMU внутри программы приводится к прежнему виду time, wx…wz, ax…az. "
        "Алгоритм Калмана тот же, что в исходной программе.",
    )
    add_para(
        doc,
        "Предсказание фильтра выполняется на каждом такте ИМУ. Полная коррекция "
        "ins::correct() — при появлении нового кадра GPS (на этой записи около 6 Гц). "
        "Дополнительно на спокойных участках работает коррекция тангажа и крена по акселерометру.",
    )

    add_heading(doc, "2. Результаты: графики", 1)

    add_heading(doc, "2.1. Калман и GPS на одних осях", 2)
    add_picture(
        doc,
        FIG / "result_comparison.png",
        "Рисунок 1. Сравнение решения Калмана и эталона. "
        "Синяя линия — программа (result.txt), оранжевая — GPS / истинные углы.",
        16.2,
    )
    add_table(
        doc,
        ["График", "Что видно"],
        [
            ["Долгота, широта", "квадрат полёта; синяя линия почти совпадает с оранжевой"],
            ["Высота", "стоянка ~124 м, набор до ~273 м; в конце Калман чуть выше GPS"],
            ["Vn, Vh, Ve", "скорости север / вертикаль / восток; эталон ступеньками, на разворотах выбросы"],
            ["Курс", "на прямых участках близко к эталону; на углах квадрата GPS ступенька, Калман запаздывает"],
            ["Тангаж, крен", "Калман около 0°; эталон после ~400 с скачет к ±180°"],
        ],
        [5.0, 11.5],
    )
    add_para(
        doc,
        "Подписи +4.764e1 и −1.2213e2 — общая часть широты 47.64° и долготы −122.13°, "
        "которую matplotlib вынес с оси.",
    )

    add_heading(doc, "2.2. Горизонтальная траектория", 2)
    add_picture(
        doc,
        FIG / "result_map.png",
        "Рисунок 2. Вид сверху: широта–долгота. Оранжевый контур — GPS, синий — Калман.",
        13.5,
    )
    add_para(
        doc,
        "Квадрат — маршрут летательного аппарата. Три стороны синяя линия почти совпадает "
        "с оранжевой. На восточной стороне Калман уходит внутрь квадрата (ошибка по востоку "
        "до ~5 м). Это главный качественный результат: траектория узнаваема и в целом "
        "сопровождает GPS, расхождение локальное.",
    )

    add_heading(doc, "2.3. Ошибки Калман − эталон", 2)
    add_picture(
        doc,
        FIG / "result_errors.png",
        "Рисунок 3. Ошибка решения относительно эталона. Линия у нуля — совпадение с GPS/углами.",
        16.2,
    )
    add_para(doc, "Каждый график — разность «программа минус эталон» по времени:")
    add_bullet(doc, "линия у нуля — совпадение с эталоном;")
    add_bullet(doc, "выше нуля — Калман больше эталона;")
    add_bullet(doc, "ниже нуля — Калман меньше эталона.")
    add_table(
        doc,
        ["График", "Смысл"],
        [
            ["Ошибка север / восток", "промах по горизонтали, м"],
            ["Ошибка высоты", "промах по высоте, м"],
            ["Ошибка Vn, Vh, Ve", "промах по скорости, м/с"],
            ["Ошибка курса, тангажа, крена", "промах по углам, град"],
        ],
        [7.0, 9.5],
    )
    add_para(
        doc,
        "До ~400 с ошибки положения малы (доли метра). На облёте квадрата, особенно "
        "после ~500 с, ошибка севера доходит до −4 м, востока — до 5 м, высоты — до 2 м. "
        "Скорости портятся на тех же разворотах (выбросы Vn до 4 м/с). "
        "Ошибка курса на прямых участках около нуля, на углах — всплески до 170° из‑за "
        "запаздывания относительно ступеньки эталона.",
    )
    add_para(
        doc,
        "Плоские участки ошибки тангажа и крена около ±180° — это не «переворот» БИНС. "
        "Эталон angle.data после ~400 с переходит на ±180°, а решение Калмана остаётся "
        "около 0°, что согласуется с акселерометром на горизонте.",
    )

    add_heading(doc, "2.4. Вектор состояния фильтра Калмана", 2)
    add_picture(
        doc,
        FIG / "kalman_state_plots.png",
        "Рисунок 4. Вектор ошибок фильтра Калмана x после коррекции по СНС. "
        "dN, dE — ошибки широты и долготы, переведённые в метры.",
        16.2,
    )
    add_para(
        doc,
        "Это не ошибка «Калман − GPS» с рисунка 3, а внутренний вектор x фильтра: "
        "на сколько фильтр предлагает поправить номинал в момент кадра СНС. "
        "После коррекции ошибки вливаются в номинал и обнуляются, поэтому масштаб "
        "малый: dN и dE в пределах метра, высота — доли метра, ba и bg — единицы "
        "10⁻⁴…10⁻³. Всплески около 500–600 с совпадают с разворотами квадрата.",
    )
    add_para(
        doc,
        "В начале записи dpsi ≈ −17°: это расхождение курса выставки (−16.7°) "
        "с эталоном. Фильтр сразу вливает эту поправку, дальше ошибка курса в x "
        "держится около нуля до манёвров.",
    )

    add_heading(doc, "3. Численные ошибки (RMSE)", 1)
    add_para(
        doc,
        "RMSE — среднеквадратичная ошибка Калман − эталон за весь полёт "
        "(по меткам времени GPS). Чем меньше, тем ближе решение к эталону. "
        "«Макс. |e|» — наибольший модуль ошибки на записи.",
    )
    add_table(
        doc,
        ["Величина", "RMSE", "Макс. |ошибка|"],
        [
            ["Север, м", "0.56", "3.89"],
            ["Восток, м", "0.78", "4.96"],
            ["Горизонт, м", "0.96", "5.86"],
            ["Высота, м", "0.56", "2.00"],
            ["Vn, м/с", "0.26", "3.75"],
            ["Vh, м/с", "0.10", "0.77"],
            ["Ve, м/с", "0.27", "2.65"],
            ["Курс, град", "20.76", "171.04"],
            ["Тангаж, град", "85.47", "180.68"],
            ["Крен, град", "86.02", "184.49"],
        ],
        [6.0, 5.0, 5.5],
    )
    add_para(
        doc,
        "Положение. RMSE по горизонту 0.96 м, по высоте 0.56 м. Максимумы 2–6 м "
        "приходятся на восточную сторону квадрата и развороты после ~500 с. "
        "Для данной записи это устойчивое сопровождение GPS.",
    )
    add_para(
        doc,
        "Скорость. Вертикальная компонента спокойная (RMSE 0.10 м/с). "
        "Горизонтальные Vn и Ve имеют выбросы 2.5–4 м/с на разворотах: эталон "
        "меняется ступенькой, БИНС интегрирует ИМУ непрерывно.",
    )
    add_para(
        doc,
        "Курс. RMSE 21° в основном из‑за запаздывания на углах квадрата и из‑за "
        "начальной выставки. На прямолинейных участках синяя и оранжевая линии близки.",
    )
    add_para(
        doc,
        "Тангаж и крен. RMSE ~85° и максимум ~180° нельзя читать как плохую выставку. "
        "На графике сравнения эталон после ~400 с уходит на ±180°, Калман остаётся около 0°. "
        "Разница почти 180° даёт огромный RMSE. Возможные причины: скачок эталона через "
        "разрез ±π, иной порядок углов Эйлера в angle.data или другая привязка СК тела. "
        "По акселерометру на стоянке крен и тангаж Калмана близки к нулю — это физически правдоподобно.",
    )

    add_heading(doc, "4. Способы уменьшить ошибки", 1)
    add_para(
        doc,
        "По графикам ошибки разной природы: часть — реальная погрешность БИНС, "
        "часть — эталон и единицы измерения. Их нужно уменьшать по отдельности.",
    )

    add_heading(doc, "4.1. Что сейчас даёт ошибку", 2)
    add_para(
        doc,
        "Положение (RMSE около 1 м, до 6 м на восточной стороне). Квадрат в целом "
        "держится. Уход на одной стороне — накопление после разворотов: курс запаздывает, "
        "скорость на повороте врёт, интеграл уводит траекторию.",
    )
    add_para(
        doc,
        "Скорости (выбросы 2–4 м/с). GPS меняется ступенькой, БИНС — непрерывно. "
        "На угле квадрата инновация скорости большая, фильтр дёргает решение.",
    )
    add_para(
        doc,
        "Курс (RMSE около 21°, пики до 170°). Выставка дала −16.7° вместо примерно 0°. "
        "На прямых фильтр подтягивает курс к angle.data. На развороте эталон скачет, "
        "гироскоп крутит плавно — отставание.",
    )
    add_para(
        doc,
        "Тангаж и крен (RMSE около 85°). Это почти наверняка не «самолёт перевернулся». "
        "Калман держит около 0° (так и должен горизонт по акселерометру). Эталон после "
        "~400 с прыгает на ±180°. Пока эталон не поправить, RMSE углов не уменьшится "
        "никаким фильтром Калмана.",
    )

    add_heading(doc, "4.2. Что делать, по убыванию пользы", 2)
    add_para(
        doc,
        "1. Починить эталон углов (самое дешёвое для RMSE тангажа и крена). "
        "Разворачивать углы (unwrap), чтобы не было скачка ±180°. Проверить порядок "
        "Эйлера и систему координат в angle.data относительно тела (X вперёд, Y вверх, "
        "Z вправо). Не кормить фильтр pitch/roll из angle.data, пока конвенция не совпадает: "
        "у нас они уже почти выключены огромным R, но курс из того же файла всё ещё измерение.",
    )
    add_para(
        doc,
        "2. Не тянуть курс к эталону на манёвре. Сейчас ins::correct() берёт курс при каждом "
        "новом GPS. На развороте лучше поднимать R курса, когда |ω| большое, или вовсе не "
        "измерять курс, пока гироскоп «шумный»: курс оставлять гироскопу, GPS — только "
        "координаты и скорости. Иначе фильтр пытается догнать ступеньку и портит и курс, "
        "и скорости.",
    )
    add_para(
        doc,
        "3. Адаптивный шум измерений GPS. Сейчас R постоянный (около 5 м, 0.1 м/с, 1°). "
        "На стоянке GPS можно доверять сильнее, на развороте — слабее. Тогда не будет "
        "острых пиков Vn и Ve.",
    )
    add_para(
        doc,
        "4. Проверить оси скоростей в gps.data. Файл даёт vx, vy, vz, а программа считает "
        "это Vn, Vh, Ve. Если в датасете ENU (восток–север–верх) или скорости в связанной "
        "СК, фильтр всё время «исправляет» не ту компоненту — типичный источник ухода "
        "на одной стороне квадрата.",
    )
    add_para(
        doc,
        "5. Рычаг антенны GPS и синхронизация. ИМУ и GPS разнесены, время у них разное. "
        "Нулевое удержание GPS (держим прошлый кадр до нового) даёт ступеньки. "
        "Интерполяция GPS на момент ИМУ и учёт lever arm уменьшают и скорость, и метры на карте.",
    )
    add_para(
        doc,
        "6. Выставка. Оценить не только ba, но и bg; дольше стоять или чуть покачать. "
        "Стартовый dψ ≈ −17° сейчас сразу выливается в фильтр — лучше начать ближе "
        "к истинному курсу.",
    )
    add_para(
        doc,
        "7. Расширить модель, если пунктов 1–6 мало. 15 состояний знают только смещения ba и bg. "
        "Нет масштабных коэффициентов гироскопа/акселерометра и перекоса осей. "
        "На длинном квадрате это даёт медленный уход по одной стороне.",
    )
    add_para(
        doc,
        "8. Не путать «ошибку фильтра x» с «ошибкой относительно GPS». На графике состояния "
        "вектор x после коррекции обнуляется, поэтому там доли метра. Реальная ошибка "
        "траектории — на графике «Калман − GPS». Уменьшать нужно её, а не амплитуду x.",
    )
    add_para(
        doc,
        "Практически для этой записи порядок такой: сначала unwrap и система координат углов, "
        "затем не корректировать курс на повороте, затем проверить оси vx, vy, vz. "
        "Эти три шага бьют в то, что видно на рисунках: ±180° на тангаже и крене, "
        "пики курса и скоростей и изгиб восточной стороны квадрата.",
    )

    add_heading(doc, "5. Выводы", 1)
    add_bullet(
        doc,
        "Данные с Google Drive (https://drive.google.com/file/d/1pox9GSTWwSeGekAxN1-u2y8ycIR2mtac/view?usp=sharing) "
        "обработаны программой imitator в каталоге newformat.",
    )
    add_bullet(
        doc,
        "ИМУ — imu.data, эталон координат и скоростей — gps.data, истинные углы — angle.data.",
    )
    add_bullet(
        doc,
        "Горизонтальная траектория Калмана повторяет квадрат GPS. RMSE по горизонту 0.96 м, "
        "по высоте 0.56 м. Локальный уход — на восточной стороне квадрата (до ~5 м).",
    )
    add_bullet(
        doc,
        "Коррекция по СНС выполняется на каждом новом кадре GPS (~6 Гц на этой записи); "
        "предсказание фильтра — на каждом такте ИМУ.",
    )
    add_bullet(
        doc,
        "Основные расхождения — запаздывание курса и скоростей на разворотах. "
        "Большие RMSE тангажа и крена связаны со скачком эталона к ±180°, а не с тем, "
        "что решение БИНС «перевернулось».",
    )
    add_bullet(
        doc,
        "Чтобы уменьшить ошибки: сначала поправить эталон углов (unwrap, система координат), "
        "не корректировать курс на манёвре, проверить оси скоростей gps.data. "
        "Далее — адаптивный R, синхронизация GPS/ИМУ, оценка bg на выставке.",
    )

    try:
        doc.save(OUT)
        print(f"Сохранено: {OUT}")
    except PermissionError:
        doc.save(OUT_ALT)
        print(f"Исходный файл открыт, сохранено: {OUT_ALT}")


if __name__ == "__main__":
    build()
