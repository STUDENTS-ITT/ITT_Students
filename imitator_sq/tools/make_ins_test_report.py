"""Отчёт по тесту автономной БИНС (столбцы *_ins в kalman15_line2.txt)."""

from __future__ import annotations

import sys
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


def load(path: Path) -> tuple[list[str], np.ndarray]:
    lines = path.read_text(encoding="utf-8", errors="replace").splitlines()
    header = lines[0].split()
    rows = []
    for line in lines[1:]:
        parts = line.split()
        if len(parts) < len(header):
            continue
        try:
            rows.append([float(x) for x in parts[1:]])
        except ValueError:
            continue
    names = header[1:]
    return names, np.asarray(rows, dtype=float)


def col(names: list[str], data: np.ndarray, name: str) -> np.ndarray:
    return data[:, names.index(name)]


def rmse(v: np.ndarray) -> float:
    return float(np.sqrt(np.mean(np.square(v))))


def gps_rows(names: list[str], data: np.ndarray) -> np.ndarray:
    keys = ("lon_sns", "lat_sns", "alt_sns", "vn_sns", "vh_sns", "ve_sns")
    block = np.column_stack([col(names, data, k) for k in keys])
    changed = np.ones(len(data), dtype=bool)
    if len(data) > 1:
        changed[1:] = np.any(np.abs(np.diff(block, axis=0)) > 1e-12, axis=1)
    return data[changed]


def set_run_font(run, size=12, bold=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size=16 if level == 1 else 14, bold=True)


def add_para(doc, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(8)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.first_line_indent = Cm(1.25)
    run = p.add_run(text)
    set_run_font(run)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def shade_header(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "D9E2F3")
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def add_picture(doc, path: Path, caption: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.add_run().add_picture(str(path), width=Cm(16.0))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.first_line_indent = Cm(0)
    run = cap.add_run(caption)
    set_run_font(run, size=11)
    run.italic = True


def main() -> int:
    root = Path(__file__).resolve().parent
    names, data = load(root / "kalman15_line2.txt")
    gps = gps_rows(names, data)
    t = col(names, gps, "time")

    err_names = [
        ("Север, м", "dN_ins"),
        ("Восток, м", "dE_ins"),
        ("Высота, м", "dh_ins"),
        ("Vn, м/с", "dVn_ins"),
        ("Vh, м/с", "dVh_ins"),
        ("Ve, м/с", "dVe_ins"),
        ("Курс, °", "dhdg_ins"),
        ("Тангаж, °", "dpitch_ins"),
        ("Крен, °", "droll_ins"),
    ]
    horiz = np.hypot(col(names, gps, "dN_ins"), col(names, gps, "dE_ins"))
    rows = []
    series = {}
    for title, key in err_names:
        v = col(names, gps, key)
        series[key] = v
        rows.append((title, f"{rmse(v):.4f}", f"{float(np.mean(v)):.4f}", f"{float(np.max(np.abs(v))):.4f}"))
    rows.insert(2, ("Горизонт, м", f"{rmse(horiz):.4f}", f"{float(np.mean(horiz)):.4f}", f"{float(np.max(horiz)):.4f}"))

    fig_dir = root / "report_ins_figures"
    fig_dir.mkdir(parents=True, exist_ok=True)

    fig, axes = plt.subplots(3, 3, figsize=(16, 11), sharex=True)
    fig.suptitle("Автономная БИНС и СНС", fontsize=14)
    pairs = [
        (axes[0, 0], "lon_ins", "lon_sns", "Долгота, °"),
        (axes[0, 1], "lat_ins", "lat_sns", "Широта, °"),
        (axes[0, 2], "alt_ins", "alt_sns", "Высота, м"),
        (axes[1, 0], "vn_ins", "vn_sns", "Vn, м/с"),
        (axes[1, 1], "vh_ins", "vh_sns", "Vh, м/с"),
        (axes[1, 2], "ve_ins", "ve_sns", "Ve, м/с"),
        (axes[2, 0], "heading_ins", "heading_sns", "Курс, °"),
        (axes[2, 1], "pitch_ins", "pitch_sns", "Тангаж, °"),
        (axes[2, 2], "roll_ins", "roll_sns", "Крен, °"),
    ]
    for ax, a, b, ylab in pairs:
        ax.plot(t, col(names, gps, a), label="Автономная БИНС", color="#0072bd", linewidth=1.1)
        ax.plot(t, col(names, gps, b), label="СНС", color="#d95319", linewidth=1.0, alpha=0.9)
        ax.set_ylabel(ylab)
        ax.grid(True, alpha=0.3)
        ax.legend(fontsize=8)
    for ax in axes[2]:
        ax.set_xlabel("Время, с")
    fig.tight_layout()
    fig.savefig(fig_dir / "ins_comparison.png", dpi=150)
    plt.close(fig)

    fig, axes = plt.subplots(3, 3, figsize=(16, 11), sharex=True)
    fig.suptitle("Ошибка автономной БИНС − СНС", fontsize=14)
    panels = [
        (axes[0, 0], series["dN_ins"], "Δ север, м"),
        (axes[0, 1], series["dE_ins"], "Δ восток, м"),
        (axes[0, 2], series["dh_ins"], "Δ высота, м"),
        (axes[1, 0], series["dVn_ins"], "Δ Vn, м/с"),
        (axes[1, 1], series["dVh_ins"], "Δ Vh, м/с"),
        (axes[1, 2], series["dVe_ins"], "Δ Ve, м/с"),
        (axes[2, 0], series["dhdg_ins"], "Δ курс, °"),
        (axes[2, 1], series["dpitch_ins"], "Δ тангаж, °"),
        (axes[2, 2], series["droll_ins"], "Δ крен, °"),
    ]
    for ax, y, ylab in panels:
        ax.plot(t, y, color="#c0392b", linewidth=1.0)
        ax.axhline(0.0, color="0.45", linewidth=0.8)
        ax.set_ylabel(ylab)
        ax.grid(True, alpha=0.3)
    for ax in axes[2]:
        ax.set_xlabel("Время, с")
    fig.tight_layout()
    fig.savefig(fig_dir / "ins_errors.png", dpi=150)
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(8, 8))
    ax.plot(col(names, gps, "lon_sns"), col(names, gps, "lat_sns"), label="СНС", color="#d95319")
    ax.plot(col(names, gps, "lon_ins"), col(names, gps, "lat_ins"), label="Автономная БИНС", color="#0072bd")
    ax.set_xlabel("Долгота, °")
    ax.set_ylabel("Широта, °")
    ax.set_title("Траектория автономной БИНС")
    ax.legend()
    ax.grid(True, alpha=0.3)
    ax.set_aspect("equal", adjustable="datalim")
    fig.tight_layout()
    fig.savefig(fig_dir / "ins_map.png", dpi=150)
    plt.close(fig)

    doc = Document()
    add_heading(doc, "Тест алгоритма автономной БИНС")
    add_para(
        doc,
        f"Параллельно с фильтром Калмана на тех же отсчётах ИМУ считалась автономная БИНС. "
        f"Начальные координаты, углы и смещения акселерометра взяты из выставки и дальше не правились. "
        f"Калман, коррекция тангажа и крена по акселерометру и обратная связь по скорости в этот канал не входят. "
        f"Сравнение с СНС на обновлениях эталона, интервал {t[0]:.1f}–{t[-1]:.1f} с, точек {len(gps)}. "
        f"Ошибка — автономная БИНС минус СНС.",
    )
    add_heading(doc, "1. Ошибки автономной БИНС", 2)
    table = doc.add_table(rows=1 + len(rows), cols=4)
    table.style = "Table Grid"
    for i, header in enumerate(("Величина", "RMSE", "Среднее", "max |e|")):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(header)
        set_run_font(run, size=11, bold=True)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        shade_header(cell)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            cell.text = ""
            run = cell.paragraphs[0].add_run(val)
            set_run_font(run, size=11)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 else WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

    add_heading(doc, "2. Сравнение со СНС", 2)
    add_picture(doc, fig_dir / "ins_comparison.png", "Рис. 1 — автономная БИНС и СНС")
    add_picture(doc, fig_dir / "ins_map.png", "Рис. 2 — траектория в плане")
    add_heading(doc, "3. Графики ошибок", 2)
    add_picture(doc, fig_dir / "ins_errors.png", "Рис. 3 — ошибка автономной БИНС − СНС")
    add_heading(doc, "4. Что показывает тест", 2)
    add_para(
        doc,
        "Без коррекции ошибки растут со временем: скорость уходит из-за остаточного смещения акселерометра "
        "и погрешности ориентации, координаты — как интеграл этой скорости. "
        "Курс, тангаж и крен держит только гироскоп, поэтому к концу пролёта угловые ошибки больше, "
        "чем у решения с фильтром Калмана. Столбцы lon_ins … ve_ins — само решение, "
        "dN_ins … droll_ins — его ошибка относительно СНС.",
    )
    out = root / "Отчет_тест_ИНС.docx"
    doc.save(out)
    print(f"точек {len(gps)}  {t[0]:.1f}–{t[-1]:.1f} с")
    for row in rows:
        print(f"{row[0]:<14} RMSE={row[1]:>12}  mean={row[2]:>12}  max={row[3]:>12}")
    print(out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
