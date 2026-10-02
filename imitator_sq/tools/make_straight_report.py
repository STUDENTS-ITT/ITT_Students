"""Отчёт Word: прямолинейный участок — Калман+БИНС vs СНС."""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# reuse column layout from plot_trajectory
COL_TIME = 0
COL_LON, COL_LAT, COL_ALT = 1, 2, 3
COL_HDG, COL_PITCH, COL_ROLL = 4, 5, 6
COL_VN, COL_VH, COL_VE = 7, 8, 9
COL_LON_GPS, COL_LAT_GPS, COL_ALT_GPS = 10, 11, 12
COL_HDG_GPS, COL_PITCH_GPS, COL_ROLL_GPS = 13, 14, 15
COL_VN_GPS, COL_VH_GPS, COL_VE_GPS = 16, 17, 18
N_COMBINED = 19
R_EARTH = 6371000.0


def load(path: Path) -> np.ndarray:
    rows: list[list[float]] = []
    with path.open(encoding="utf-8", errors="replace") as f:
        for line in f:
            parts = line.split()
            if not parts:
                continue
            start = 1 if parts[0].count(":") == 2 else 0
            try:
                rows.append([float(x) for x in parts[start:]])
            except ValueError:
                continue
    if not rows:
        raise ValueError(f"Нет данных в {path}")
    return np.asarray(rows, dtype=float)


def gps_update_mask(data: np.ndarray) -> np.ndarray:
    gps = data[:, COL_LON_GPS : COL_VE_GPS + 1]
    changed = np.ones(len(data), dtype=bool)
    if len(data) > 1:
        changed[1:] = np.any(np.abs(gps[1:] - gps[:-1]) > 1e-12, axis=1)
    return changed


def wrap_deg(angle: np.ndarray) -> np.ndarray:
    return (angle + 180.0) % 360.0 - 180.0


def rmse(values: np.ndarray) -> float:
    return float(np.sqrt(np.mean(np.square(values))))


def errors_from_gps_rows(data: np.ndarray) -> dict[str, np.ndarray]:
    dlat = np.deg2rad(data[:, COL_LAT] - data[:, COL_LAT_GPS])
    dlon = np.deg2rad(data[:, COL_LON] - data[:, COL_LON_GPS])
    north = dlat * R_EARTH
    east = dlon * R_EARTH * np.cos(np.deg2rad(data[:, COL_LAT_GPS]))
    return {
        "north_m": north,
        "east_m": east,
        "horiz_m": np.hypot(north, east),
        "alt": data[:, COL_ALT] - data[:, COL_ALT_GPS],
        "vn": data[:, COL_VN] - data[:, COL_VN_GPS],
        "vh": data[:, COL_VH] - data[:, COL_VH_GPS],
        "ve": data[:, COL_VE] - data[:, COL_VE_GPS],
        "hdg": wrap_deg(data[:, COL_HDG] - data[:, COL_HDG_GPS]),
        "pitch": data[:, COL_PITCH] - data[:, COL_PITCH_GPS],
        "roll": data[:, COL_ROLL] - data[:, COL_ROLL_GPS],
    }


def find_straight_segment(gps_rows: np.ndarray) -> tuple[float, float]:
    """Интервал [t0, t1] по отсчётам СНС: |Δкурс|<2° между шагами, |V|>0.5 м/с."""
    if len(gps_rows) < 5:
        t = gps_rows[:, COL_TIME]
        return float(t[0]), float(t[-1])

    tg = gps_rows[:, COL_TIME]
    hdg = gps_rows[:, COL_HDG_GPS]
    spd = np.hypot(gps_rows[:, COL_VN_GPS], gps_rows[:, COL_VE_GPS])

    dh = np.abs(wrap_deg(hdg[1:] - hdg[:-1]))
    ok = (dh < 2.0) & (spd[1:] > 0.5)

    best_i, best_j, best_len = 0, 1, 0
    i = 0
    while i < len(ok):
        if not ok[i]:
            i += 1
            continue
        j = i
        while j < len(ok) and ok[j]:
            j += 1
        if j - i > best_len:
            best_len = j - i
            best_i, best_j = i, j
        i = j

    if best_len >= 10:
        t0 = float(tg[best_i])
        t1 = float(tg[min(best_j, len(tg) - 1)])
        if t1 - t0 >= 5.0:
            return t0, t1

    # запасной вариант: участок с минимальным std курса среди окон ≥30 с
    t0_all, t1_all = float(tg[0]), float(tg[-1])
    win = 30.0
    step = 5.0
    best_std = 1e9
    best_t0, best_t1 = t0_all, min(t0_all + win, t1_all)
    t = t0_all
    while t + win <= t1_all:
        m = (tg >= t) & (tg <= t + win)
        if np.sum(m) >= 8:
            s = float(np.std(hdg[m]))
            if s < best_std:
                best_std = s
                best_t0, best_t1 = t, t + win
        t += step
    return best_t0, best_t1


def configure_style() -> None:
    plt.rcParams.update(
        {
            "figure.facecolor": "white",
            "axes.grid": True,
            "grid.alpha": 0.35,
            "font.size": 10,
            "lines.linewidth": 1.0,
        }
    )


def plot_comparison(gps: np.ndarray, out_path: Path, title: str) -> None:
    t = gps[:, COL_TIME]
    fig, axes = plt.subplots(3, 3, figsize=(16, 11), sharex=True)
    fig.suptitle(title, fontsize=14)
    panels = [
        (axes[0, 0], COL_LON, COL_LON_GPS, "Долгота, °"),
        (axes[0, 1], COL_LAT, COL_LAT_GPS, "Широта, °"),
        (axes[0, 2], COL_ALT, COL_ALT_GPS, "Высота, м"),
        (axes[1, 0], COL_VN, COL_VN_GPS, "Vn, м/с"),
        (axes[1, 1], COL_VH, COL_VH_GPS, "Vh, м/с"),
        (axes[1, 2], COL_VE, COL_VE_GPS, "Ve, м/с"),
        (axes[2, 0], COL_HDG, COL_HDG_GPS, "Курс, °"),
        (axes[2, 1], COL_PITCH, COL_PITCH_GPS, "Тангаж, °"),
        (axes[2, 2], COL_ROLL, COL_ROLL_GPS, "Крен, °"),
    ]
    kc, gc = "#0072bd", "#d95319"
    for ax, ck, cg, ylab in panels:
        ax.plot(t, gps[:, ck], label="БИНС+Калман", color=kc, linewidth=1.2)
        ax.plot(t, gps[:, cg], label="СНС (эталон)", color=gc, linewidth=1.0, alpha=0.9)
        ax.set_ylabel(ylab)
        ax.legend(fontsize=8)
    for ax in axes[2]:
        ax.set_xlabel("Время, с")
    fig.tight_layout()
    out_path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(out_path, dpi=150)
    plt.close(fig)


def plot_errors(t: np.ndarray, err: dict[str, np.ndarray], out_path: Path, title: str) -> None:
    fig, axes = plt.subplots(3, 3, figsize=(16, 11), sharex=True)
    fig.suptitle(title, fontsize=14)
    panels = [
        (axes[0, 0], err["north_m"], "Δ север, м"),
        (axes[0, 1], err["east_m"], "Δ восток, м"),
        (axes[0, 2], err["alt"], "Δ высота, м"),
        (axes[1, 0], err["vn"], "Δ Vn, м/с"),
        (axes[1, 1], err["vh"], "Δ Vh, м/с"),
        (axes[1, 2], err["ve"], "Δ Ve, м/с"),
        (axes[2, 0], err["hdg"], "Δ курс, °"),
        (axes[2, 1], err["pitch"], "Δ тангаж, °"),
        (axes[2, 2], err["roll"], "Δ крен, °"),
    ]
    for ax, y, ylab in panels:
        ax.plot(t, y, color="#c0392b", linewidth=1.0)
        ax.axhline(0.0, color="0.45", linewidth=0.8)
        ax.set_ylabel(ylab)
    for ax in axes[2]:
        ax.set_xlabel("Время, с")
    fig.tight_layout()
    out_path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(out_path, dpi=150)
    plt.close(fig)


def plot_velocity_sources(gps: np.ndarray, idx: dict[str, int], out_path: Path, title: str) -> None:
    t = gps[:, COL_TIME]
    fig, axes = plt.subplots(3, 1, figsize=(12, 9), sharex=True)
    fig.suptitle(title, fontsize=14)
    series = [
        ("Калман", "#0072bd", ("vn", "vh", "ve")),
        ("СНС", "#d95319", ("vn_sns", "vh_sns", "ve_sns")),
        ("ИНС", "#2ca02c", ("vn_ins", "vh_ins", "ve_ins")),
        ("V_free", "#7b2d8e", ("vn_free", "vh_free", "ve_free")),
    ]
    ylabels = ("Vn, м/с", "Vh, м/с", "Ve, м/с")
    for ax, ylab, k in zip(axes, ylabels, range(3)):
        for label, color, keys in series:
            ax.plot(t, gps[:, idx[keys[k]]], label=label, color=color, linewidth=1.15)
        ax.set_ylabel(ylab)
        ax.grid(True, alpha=0.3)
        ax.legend(fontsize=8, ncol=4, loc="best")
    axes[-1].set_xlabel("Время, с")
    fig.tight_layout()
    out_path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(out_path, dpi=150)
    plt.close(fig)


def plot_map(gps: np.ndarray, out_path: Path, title: str) -> None:
    fig, ax = plt.subplots(figsize=(8, 8))
    ax.plot(
        gps[:, COL_LON_GPS],
        gps[:, COL_LAT_GPS],
        label="СНС",
        color="#d95319",
        linewidth=1.5,
    )
    ax.plot(
        gps[:, COL_LON],
        gps[:, COL_LAT],
        label="БИНС+Калман",
        color="#0072bd",
        linewidth=1.2,
    )
    ax.set_xlabel("Долгота, °")
    ax.set_ylabel("Широта, °")
    ax.set_title(title)
    ax.legend()
    ax.set_aspect("equal", adjustable="datalim")
    fig.tight_layout()
    out_path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(out_path, dpi=150)
    plt.close(fig)


def set_run_font(run, size=12, bold=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size=16 if level == 1 else 14, bold=True)


def add_para(doc, text, *, first_line=True):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(8)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if first_line:
        pf.first_line_indent = Cm(1.25)
    run = p.add_run(text)
    set_run_font(run)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def add_picture(doc, path: Path, caption: str, width_cm=16.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.first_line_indent = Cm(0)
    r = cap.add_run(caption)
    set_run_font(r, size=11)
    r.italic = True


def stats_table_rows(err: dict[str, np.ndarray]) -> list[tuple[str, str, str]]:
    items = [
        ("Север, м", err["north_m"]),
        ("Восток, м", err["east_m"]),
        ("Горизонт, м", err["horiz_m"]),
        ("Высота, м", err["alt"]),
        ("Vn, м/с", err["vn"]),
        ("Vh, м/с", err["vh"]),
        ("Ve, м/с", err["ve"]),
        ("Курс, °", err["hdg"]),
        ("Тангаж, °", err["pitch"]),
        ("Крен, °", err["roll"]),
    ]
    rows = []
    for name, v in items:
        rows.append((name, f"{rmse(v):.4f}", f"{float(np.max(np.abs(v))):.4f}"))
    return rows


def build_docx(
    out_doc: Path,
    fig_dir: Path,
    t0: float,
    t1: float,
    stats: list[tuple[str, str, str]],
    source: Path,
    n_gps: int,
    mean_speed: float,
    hdg_std: float,
) -> None:
    doc = Document()
    add_heading(doc, "Отчёт по результатам навигации (имитатор, прямолинейный участок)")
    add_para(
        doc,
        f"Источник данных: {source.name}. Анализ выполнен по отсчётам СНС (моменты обновления "
        f"эталона в логе). Автоматически выделен прямолинейный участок: {t0:.1f}–{t1:.1f} с "
        f"({n_gps} точек, средняя горизонтальная скорость эталона {mean_speed:.2f} м/с, "
        f"σ курса по СНС {hdg_std:.4f}°). Сравниваются БИНС с фильтром Калмана 15-го порядка "
        f"и эталон (gps.dat + angle.dat).",
    )

    add_heading(doc, "1. Численные показатели ошибки", 2)
    table = doc.add_table(rows=1 + len(stats), cols=3)
    table.style = "Table Grid"
    hdr = ("Величина", "RMSE", "max |e|")
    for i, h in enumerate(hdr):
        table.rows[0].cells[i].text = h
    for ri, row in enumerate(stats):
        for ci, val in enumerate(row):
            table.rows[ri + 1].cells[ci].text = val
    doc.add_paragraph()

    add_heading(doc, "2. Сравнение БИНС+Калман и СНС", 2)
    add_picture(
        doc,
        fig_dir / "straight_comparison.png",
        "Рис. 1 — координаты, скорости и углы на одних осях",
    )

    add_heading(doc, "3. Ошибки (Калман − СНС)", 2)
    add_picture(
        doc,
        fig_dir / "straight_errors.png",
        "Рис. 2 — ошибки по каналам на прямолинейном участке",
    )

    add_heading(doc, "4. Горизонтальная траектория", 2)
    add_picture(
        doc,
        fig_dir / "straight_map.png",
        "Рис. 3 — наложение траекторий в плане",
    )

    add_heading(doc, "5. Скорости: Калман, СНС, ИНС и V_free", 2)
    add_picture(
        doc,
        fig_dir / "straight_velocities.png",
        "Рис. 4 — северная, вертикальная и восточная скорости. "
        "Калман — замкнутая скорость БИНС после фильтра. "
        "СНС — эталон из gps.dat. "
        "ИНС — автономное счисление без Калмана. "
        "V_free — тот же прирост скорости, что у Калмана, без вычитания поправки фильтра.",
    )

    hdg = next(s for s in stats if s[0].startswith("Курс"))
    north = next(s for s in stats if s[0].startswith("Север"))
    east = next(s for s in stats if s[0].startswith("Восток"))
    add_heading(doc, "6. Краткие выводы", 2)
    add_para(
        doc,
        f"На прямолинейном участке ошибка Калмана относительно СНС: север RMSE {north[1]} м "
        f"(макс. {north[2]} м), восток RMSE {east[1]} м. Курс: RMSE {hdg[1]}°, максимум {hdg[2]}°. "
        f"На графике скоростей Калман держится у СНС. ИНС без коррекции уходит. "
        f"V_free повторяет приращения Калмана, но не получает вычитание ошибки скорости, "
        f"поэтому тоже расходится с эталоном.",
    )

    out_doc.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_doc)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "path",
        nargs="?",
        default=None,
        help="kalman15_line2.txt",
    )
    parser.add_argument(
        "--t0",
        type=float,
        default=None,
        help="Начало участка, с (если не задано — авто)",
    )
    parser.add_argument(
        "--t1",
        type=float,
        default=None,
        help="Конец участка, с",
    )
    args = parser.parse_args()

    script_dir = Path(__file__).resolve().parent
    data_path = Path(args.path).resolve() if args.path else script_dir / "kalman15_line2.txt"
    fig_dir = script_dir / "report_straight_figures"
    out_doc = script_dir / "Отчет_прямолинейный_участок.docx"

    configure_style()
    data = load(data_path)
    if data.shape[1] < N_COMBINED:
        raise SystemExit(f"Ожидалось ≥{N_COMBINED} столбцов, получено {data.shape[1]}")

    gps_all = data[gps_update_mask(data)]
    if args.t0 is not None and args.t1 is not None:
        t0, t1 = args.t0, args.t1
    else:
        t0, t1 = find_straight_segment(gps_all)

    mask = (gps_all[:, COL_TIME] >= t0) & (gps_all[:, COL_TIME] <= t1)
    gps = gps_all[mask]
    if len(gps) < 3:
        raise SystemExit("Слишком мало точек на выбранном участке")

    err = errors_from_gps_rows(gps)
    stats = stats_table_rows(err)

    title_suffix = f"участок {t0:.1f}–{t1:.1f} с"
    plot_comparison(
        gps,
        fig_dir / "straight_comparison.png",
        f"БИНС+Калман и СНС — {title_suffix}",
    )
    plot_errors(
        gps[:, COL_TIME],
        err,
        fig_dir / "straight_errors.png",
        f"Ошибка (Калман − СНС) — {title_suffix}",
    )
    plot_map(
        gps,
        fig_dir / "straight_map.png",
        f"Траектория в плане — {title_suffix}",
    )
    header = data_path.read_text(encoding="utf-8", errors="replace").splitlines()[0].split()[1:]
    idx = {name: i for i, name in enumerate(header)}
    plot_velocity_sources(
        gps,
        idx,
        fig_dir / "straight_velocities.png",
        f"Скорости Калман, СНС, ИНС и V_free — {title_suffix}",
    )

    mean_speed = float(
        np.mean(np.hypot(gps[:, COL_VN_GPS], gps[:, COL_VE_GPS]))
    )
    hdg_std = float(np.std(gps[:, COL_HDG_GPS]))
    build_docx(
        out_doc, fig_dir, t0, t1, stats, data_path, len(gps), mean_speed, hdg_std
    )

    print(f"Участок: {t0:.2f} – {t1:.2f} с, точек СНС: {len(gps)}")
    print("RMSE (Калман − СНС):")
    for name, rm, mx in stats:
        print(f"  {name:<14} RMSE={rm:>10}  max|e|={mx}")
    print(f"Графики: {fig_dir}")
    print(f"Отчёт: {out_doc}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
