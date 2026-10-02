"""Отчёт Word: БИНС+Калман и СНС по kalman15_line2.txt."""

from __future__ import annotations

import sys
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

COL_TIME = 0
COL_LON, COL_LAT, COL_ALT = 1, 2, 3
COL_HDG, COL_PITCH, COL_ROLL = 4, 5, 6
COL_VN, COL_VH, COL_VE = 7, 8, 9
COL_LON_GPS, COL_LAT_GPS, COL_ALT_GPS = 10, 11, 12
COL_HDG_GPS, COL_PITCH_GPS, COL_ROLL_GPS = 13, 14, 15
COL_VN_GPS, COL_VH_GPS, COL_VE_GPS = 16, 17, 18
R_EARTH = 6371000.0


def load(path: Path) -> np.ndarray:
    rows: list[list[float]] = []
    with path.open(encoding="utf-8", errors="replace") as f:
        for line in f:
            parts = line.split()
            if not parts:
                continue
            start = 1 if parts[0].count(":") == 2 else 0
            try:
                rows.append([float(x) for x in parts[start:]])
            except ValueError:
                continue
    if not rows:
        raise ValueError(f"Нет данных в {path}")
    return np.asarray(rows, dtype=float)


def gps_updates(data: np.ndarray) -> np.ndarray:
    gps = data[:, COL_LON_GPS : COL_VE_GPS + 1]
    changed = np.ones(len(data), dtype=bool)
    if len(data) > 1:
        changed[1:] = np.any(np.abs(gps[1:] - gps[:-1]) > 1e-12, axis=1)
    return data[changed]


def wrap_deg(angle: np.ndarray) -> np.ndarray:
    return (angle + 180.0) % 360.0 - 180.0


def rmse(values: np.ndarray) -> float:
    return float(np.sqrt(np.mean(np.square(values))))


def errors(data: np.ndarray) -> dict[str, np.ndarray]:
    dlat = np.deg2rad(data[:, COL_LAT] - data[:, COL_LAT_GPS])
    dlon = np.deg2rad(data[:, COL_LON] - data[:, COL_LON_GPS])
    north = dlat * R_EARTH
    east = dlon * R_EARTH * np.cos(np.deg2rad(data[:, COL_LAT_GPS]))
    return {
        "north_m": north,
        "east_m": east,
        "horiz_m": np.hypot(north, east),
        "alt": data[:, COL_ALT] - data[:, COL_ALT_GPS],
        "vn": data[:, COL_VN] - data[:, COL_VN_GPS],
        "vh": data[:, COL_VH] - data[:, COL_VH_GPS],
        "ve": data[:, COL_VE] - data[:, COL_VE_GPS],
        "hdg": wrap_deg(data[:, COL_HDG] - data[:, COL_HDG_GPS]),
        "pitch": data[:, COL_PITCH] - data[:, COL_PITCH_GPS],
        "roll": data[:, COL_ROLL] - data[:, COL_ROLL_GPS],
    }


def stats_rows(err: dict[str, np.ndarray]) -> list[tuple[str, str, str, str]]:
    items = [
        ("Север, м", err["north_m"]),
        ("Восток, м", err["east_m"]),
        ("Горизонт, м", err["horiz_m"]),
        ("Высота, м", err["alt"]),
        ("Vn, м/с", err["vn"]),
        ("Vh, м/с", err["vh"]),
        ("Ve, м/с", err["ve"]),
        ("Курс, °", err["hdg"]),
        ("Тангаж, °", err["pitch"]),
        ("Крен, °", err["roll"]),
    ]
    rows = []
    for name, values in items:
        rows.append(
            (
                name,
                f"{rmse(values):.4f}",
                f"{float(np.mean(values)):.4f}",
                f"{float(np.max(np.abs(values))):.4f}",
            )
        )
    return rows


def plot_comparison(gps: np.ndarray, path: Path) -> None:
    t = gps[:, COL_TIME]
    fig, axes = plt.subplots(3, 3, figsize=(16, 11), sharex=True)
    fig.suptitle("БИНС+Калман и СНС", fontsize=14)
    panels = [
        (axes[0, 0], COL_LON, COL_LON_GPS, "Долгота, °"),
        (axes[0, 1], COL_LAT, COL_LAT_GPS, "Широта, °"),
        (axes[0, 2], COL_ALT, COL_ALT_GPS, "Высота, м"),
        (axes[1, 0], COL_VN, COL_VN_GPS, "Vn, м/с"),
        (axes[1, 1], COL_VH, COL_VH_GPS, "Vh, м/с"),
        (axes[1, 2], COL_VE, COL_VE_GPS, "Ve, м/с"),
        (axes[2, 0], COL_HDG, COL_HDG_GPS, "Курс, °"),
        (axes[2, 1], COL_PITCH, COL_PITCH_GPS, "Тангаж, °"),
        (axes[2, 2], COL_ROLL, COL_ROLL_GPS, "Крен, °"),
    ]
    for ax, ck, cg, ylab in panels:
        ax.plot(t, gps[:, ck], label="БИНС+Калман", color="#0072bd", linewidth=1.2)
        ax.plot(t, gps[:, cg], label="СНС", color="#d95319", linewidth=1.0, alpha=0.9)
        ax.set_ylabel(ylab)
        ax.grid(True, alpha=0.3)
        ax.legend(fontsize=8)
    for ax in axes[2]:
        ax.set_xlabel("Время, с")
    fig.tight_layout()
    fig.savefig(path, dpi=150)
    plt.close(fig)


def plot_errors(t: np.ndarray, err: dict[str, np.ndarray], path: Path) -> None:
    fig, axes = plt.subplots(3, 3, figsize=(16, 11), sharex=True)
    fig.suptitle("Ошибка БИНС+Калман − СНС", fontsize=14)
    panels = [
        (axes[0, 0], err["north_m"], "Δ север, м"),
        (axes[0, 1], err["east_m"], "Δ восток, м"),
        (axes[0, 2], err["alt"], "Δ высота, м"),
        (axes[1, 0], err["vn"], "Δ Vn, м/с"),
        (axes[1, 1], err["vh"], "Δ Vh, м/с"),
        (axes[1, 2], err["ve"], "Δ Ve, м/с"),
        (axes[2, 0], err["hdg"], "Δ курс, °"),
        (axes[2, 1], err["pitch"], "Δ тангаж, °"),
        (axes[2, 2], err["roll"], "Δ крен, °"),
    ]
    for ax, y, ylab in panels:
        ax.plot(t, y, color="#c0392b", linewidth=1.0)
        ax.axhline(0.0, color="0.45", linewidth=0.8)
        ax.set_ylabel(ylab)
        ax.grid(True, alpha=0.3)
    for ax in axes[2]:
        ax.set_xlabel("Время, с")
    fig.tight_layout()
    fig.savefig(path, dpi=150)
    plt.close(fig)


def plot_map(gps: np.ndarray, path: Path) -> None:
    fig, ax = plt.subplots(figsize=(8, 8))
    ax.plot(gps[:, COL_LON_GPS], gps[:, COL_LAT_GPS], label="СНС", color="#d95319", linewidth=1.5)
    ax.plot(gps[:, COL_LON], gps[:, COL_LAT], label="БИНС+Калман", color="#0072bd", linewidth=1.2)
    ax.set_xlabel("Долгота, °")
    ax.set_ylabel("Широта, °")
    ax.set_title("Траектория в плане")
    ax.legend()
    ax.grid(True, alpha=0.3)
    ax.set_aspect("equal", adjustable="datalim")
    fig.tight_layout()
    fig.savefig(path, dpi=150)
    plt.close(fig)


def set_run_font(run, size=12, bold=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size=16 if level == 1 else 14, bold=True)


def add_para(doc, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(8)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.first_line_indent = Cm(1.25)
    run = p.add_run(text)
    set_run_font(run)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def shade_header(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "D9E2F3")
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def add_picture(doc, path: Path, caption: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.add_run().add_picture(str(path), width=Cm(16.0))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.first_line_indent = Cm(0)
    run = cap.add_run(caption)
    set_run_font(run, size=11)
    run.italic = True


def main() -> int:
    root = Path(__file__).resolve().parent
    data_path = root / "kalman15_line2.txt"
    fig_dir = root / "report_kalman_figures"
    fig_dir.mkdir(parents=True, exist_ok=True)
    out_doc = root / "Отчет_kalman15_line2.docx"

    data = load(data_path)
    gps = gps_updates(data)
    err = errors(gps)
    rows = stats_rows(err)
    t0 = float(gps[0, COL_TIME])
    t1 = float(gps[-1, COL_TIME])

    plot_comparison(gps, fig_dir / "comparison.png")
    plot_errors(gps[:, COL_TIME], err, fig_dir / "errors.png")
    plot_map(gps, fig_dir / "map.png")

    doc = Document()
    add_heading(doc, "Отчёт по результатам БИНС с фильтром Калмана")
    add_para(
        doc,
        f"Источник: {data_path.name}. Сравнение решения БИНС после коррекции фильтром Калмана "
        f"с эталоном СНС на моментах обновления эталона. Интервал {t0:.1f}–{t1:.1f} с, "
        f"точек СНС: {len(gps)}. Ошибка считается как БИНС+Калман минус СНС. "
        f"Курс эталона в файле — yaw из angle.dat минус 32°.",
    )

    add_heading(doc, "1. Ошибки", 2)
    table = doc.add_table(rows=1 + len(rows), cols=4)
    table.style = "Table Grid"
    for i, header in enumerate(("Величина", "RMSE", "Среднее", "max |e|")):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(header)
        set_run_font(run, size=11, bold=True)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        shade_header(cell)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            cell.text = ""
            run = cell.paragraphs[0].add_run(val)
            set_run_font(run, size=11)
            cell.paragraphs[0].alignment = (
                WD_ALIGN_PARAGRAPH.LEFT if ci == 0 else WD_ALIGN_PARAGRAPH.CENTER
            )
    doc.add_paragraph()

    add_heading(doc, "2. Сравнение БИНС+Калман и СНС", 2)
    add_picture(doc, fig_dir / "comparison.png", "Рис. 1 — координаты, скорости и углы")
    add_picture(doc, fig_dir / "map.png", "Рис. 2 — траектория в плане")

    add_heading(doc, "3. Графики ошибок", 2)
    add_picture(doc, fig_dir / "errors.png", "Рис. 3 — ошибка БИНС+Калман − СНС")

    add_heading(doc, "4. Кратко по графику", 2)
    add_para(
        doc,
        "На прямой курс и крен держатся у эталона. Расхождение курса нарастает после 290 с, "
        "когда аппарат меняет тангаж и останавливается: замер курса в фильтре на этом участке выключен, "
        "и угол интегрируется гироскопом. Ошибка тангажа больше всего на участке изменения наклона.",
    )

    doc.save(out_doc)
    print(f"Интервал {t0:.2f}–{t1:.2f} с, точек СНС {len(gps)}")
    for row in rows:
        print(f"{row[0]:<14} RMSE={row[1]:>10}  mean={row[2]:>10}  max={row[3]:>10}")
    print(f"Отчёт: {out_doc}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
